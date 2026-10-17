import os
import sys
import uuid
import re
import django
import openpyxl

# إعداد بيئة Django
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'config.settings')
django.setup()

from django.conf import settings
from django.db import models, transaction
from django.contrib.auth import get_user_model
from django.contrib.auth.hashers import make_password

from core_system.tenants.models import Tenant
from tenant_modules.centers_and_projects.models import Center
from tenant_modules.users.models import UserProfile
from tenant_modules.halaqat.models import Halaqa
from tenant_modules.students_and_parents.models import Student, StudentEnrollment

User = get_user_model()

def get_tenant_db(tenant_target='saadtrial'):
    """الحصول على قاعدة بيانات المستأجر"""
    tenant = Tenant.objects.filter(
        models.Q(subdomain__iexact=tenant_target) | 
        models.Q(db_name__iexact=tenant_target) | 
        models.Q(name__icontains=tenant_target)
    ).first()

    if not tenant:
        tenant = Tenant.objects.filter(is_active=True).first()

    if not tenant:
        print("⚠️ لم يتم العثور على أي مستأجر نشط!")
        return 'default', None

    db_name = tenant.db_name
    if db_name not in settings.DATABASES:
        master_db = settings.DATABASES['default']
        new_db_config = master_db.copy()
        new_db_config.update({
            'NAME': db_name,
            'USER': tenant.db_user or master_db.get('USER', 'manara_user'),
            'PASSWORD': tenant.db_password_hash or master_db.get('PASSWORD', 'M@nara_2026_Str0ng!'),
            'HOST': tenant.db_host or 'localhost',
            'PORT': tenant.db_port or 5432,
        })
        settings.DATABASES[db_name] = new_db_config

    return db_name, tenant


def clean_str(val):
    if val is None:
        return ""
    return str(val).strip()


def generate_teacher_username(teacher_name, index, db_name):
    # إنشاء اسم مستخدم فريد للمدرس
    base_username = f"teacher_{index}"
    if not User.objects.using(db_name).filter(username=base_username).exists():
        return base_username
    return f"teacher_{index}_{uuid.uuid4().hex[:4]}"


def import_excel_data(excel_file_path, tenant_target='saadtrial'):
    if not os.path.exists(excel_file_path):
        print(f"❌ الملف غير موجود في المسار: {excel_file_path}")
        return

    db_name, tenant = get_tenant_db(tenant_target)
    print(f"🔄 جاري قراءة الملف: {excel_file_path}")
    print(f"🏛️ المستأجر: {tenant.name if tenant else 'Default'} (قاعدة البيانات: {db_name})...")

    wb = openpyxl.load_workbook(excel_file_path, data_only=True)
    default_center = Center.objects.using(db_name).filter(is_active=True).first()

    total_sheets = len(wb.sheetnames)
    total_teachers = 0
    total_halaqat = 0
    total_students = 0

    print(f"📄 تم العثور على {total_sheets} ورقة عمل.")

    for sheet_idx, sheet_name in enumerate(wb.sheetnames, 1):
        sheet = wb[sheet_name]

        # قراءة اسم المدرس من F3
        teacher_name = clean_str(sheet['F3'].value)
        if not teacher_name:
            print(f"⏩ تجاوز الورقة [{sheet_name}]: خلية اسم المدرس F3 فارغة.")
            continue

        # تسمية الحلقة باسم المدرس
        halaqa_name = f"حلقة {teacher_name}"

        print(f"\n🔹 [{sheet_idx}/{total_sheets}] ورقة: {sheet_name} | المعلم: {teacher_name} | الحلقة: {halaqa_name}")

        with transaction.atomic(using=db_name):
            # 1. إنشاء / جلب المعلم
            user = User.objects.using(db_name).filter(first_name=teacher_name).first()
            if not user:
                username = generate_teacher_username(teacher_name, sheet_idx, db_name)
                user = User.objects.using(db_name).create(
                    username=username,
                    first_name=teacher_name,
                    password=make_password('12345678'),
                    is_active=True
                )
                total_teachers += 1

            # ملف المستخدم UserProfile
            user_profile, _ = UserProfile.objects.using(db_name).get_or_create(
                user=user,
                defaults={
                    'role': 'TEACHER',
                    'roles': ['TEACHER'],
                    'center': default_center,
                    'is_active': True,
                }
            )

            # 2. إنشاء / جلب الحلقة
            halaqa, created_halaqa = Halaqa.objects.using(db_name).get_or_create(
                name=halaqa_name,
                defaults={
                    'teacher_name': teacher_name,
                    'center': default_center,
                    'is_active': True,
                }
            )
            if created_halaqa:
                total_halaqat += 1

            # 3. قراءة الطلاب من العمود B ابتداءً من الصف 5
            sheet_students_count = 0
            row = 5
            consecutive_empty = 0

            while consecutive_empty < 3: # التوقف عند 3 خلايا فارغة متتالية
                student_name = clean_str(sheet[f'B{row}'].value)

                if not student_name:
                    consecutive_empty += 1
                    row += 1
                    continue

                consecutive_empty = 0

                # إنشاء الطالب
                student, created_student = Student.objects.using(db_name).get_or_create(
                    full_name=student_name,
                    halaqa=halaqa,
                    defaults={
                        'reached_page': 1,
                        'gender': 'M',
                    }
                )

                # ربط الطالب بالحلقة عبر Enrollment
                StudentEnrollment.objects.using(db_name).get_or_create(
                    student=student,
                    halaqa=halaqa,
                    defaults={'is_active': True}
                )

                sheet_students_count += 1
                total_students += 1
                row += 1

            print(f"   ✅ تم استيراد/تأكيد {sheet_students_count} طالب في {halaqa_name}")

    print("\n" + "="*50)
    print(f"🎉 تم استيراد البيانات بنجاح إلى قاعدة بيانات {db_name}:")
    print(f"   - المعلمون: {total_teachers} معلم جديد")
    print(f"   - الحلقات: {total_halaqat} حلقة جديدة")
    print(f"   - الطلاب: {total_students} طالب")
    print("="*50)


if __name__ == '__main__':
    EXCEL_FILE = sys.argv[1] if len(sys.argv) > 1 else 'students_data.xlsx'
    TENANT = sys.argv[2] if len(sys.argv) > 2 else 'saadtrial'
    import_excel_data(EXCEL_FILE, TENANT)
